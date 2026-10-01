import datetime
import os
import time

import numpy as np
import openpyxl
import torch
import torch.nn as nn
import torch.nn.functional as F
from torch.cuda import amp

import narma_estimate
import time_series_data_prediction as tp
import le_backward as lebw
from time_series_data_prediction import dataset_from


# Transformer/Informer/Autoformer的分段长度。
# 4 GB显存建议先使用32，仍然OOM可改成16。
ADVANCED_MODEL_SEGMENT_LENGTH = 32
ADVANCED_MODEL_NAMES = ('transformer', 'informer', 'autoformer')


class BatchNorm1dLN(nn.Module):
    """
    BatchNorm1d for input of shape (L, N, C)
    """

    def __init__(self, num_features):
        super().__init__()
        self.bn = nn.BatchNorm1d(num_features)

    def forward(self, x):
        # x: (L, N, C)
        L, N, C = x.shape
        # -> (N, C, L)
        x = x.permute(1, 2, 0)
        x = self.bn(x)
        # -> (L, N, C)
        x = x.permute(2, 0, 1)
        return x


class three_linear_net(nn.Module):
    def __init__(self, in_dim, n_hidden, out_dim, history_step=1, for_lsm=False):
        super().__init__()
        self.layer1 = nn.Sequential(
            nn.Linear(in_dim, n_hidden),
            BatchNorm1dLN(n_hidden),
            nn.LeakyReLU()
        )

        if for_lsm:
            self.layer2 = nn.Sequential(
                nn.Linear(2 * history_step * n_hidden, out_dim)
            )
        else:
            self.layer2 = nn.Sequential(
                nn.Linear(n_hidden, out_dim)
            )

    def forward(self, x):
        x = self.layer1(x)
        x = self.layer2(x)
        return x


def get_loss(dataset, device, net):
    if type(dataset) is tuple:
        input_dataset, output_dataset = dataset
        input_dataset = input_dataset.to(device)
        output_dataset = output_dataset.to(device)
        s, d = input_dataset.shape
        input_dataset = input_dataset.reshape(s, 1, d)
        outputs = net.forward(input_dataset).reshape(output_dataset.shape)
    else:
        outputs = torch.tensor([]).to(device)
        output_dataset_all = torch.tensor([]).to(device)
        for dataset_elem in dataset:
            input_dataset, output_dataset = dataset_elem
            input_dataset = input_dataset.to(device)
            output_dataset = output_dataset.to(device)
            s, d = input_dataset.shape
            input_dataset = input_dataset.reshape(s, 1, d)
            outputs_part = net.forward(input_dataset).reshape(output_dataset.shape)
            outputs = torch.cat([outputs, outputs_part], dim=0)
            output_dataset_all = torch.cat([output_dataset_all, output_dataset], dim=0)
        output_dataset = output_dataset_all.clone()

    loss = F.mse_loss(outputs, output_dataset)
    return loss, outputs, output_dataset


def global_gradient_l2_norm(net):
    squared_sum = 0.0

    for parameter in net.parameters():
        if parameter.requires_grad and parameter.grad is not None:
            squared_sum += parameter.grad.detach().pow(2).sum().item()

    return squared_sum ** 0.5


def _is_advanced_attention_model(net):
    """
    自动判断当前网络是否为Transformer、Informer或Autoformer。

    同时检查：
    1. 外层包装网络类名；
    2. 外层包装网络中的model成员类名；
    3. 网络及其子模块完整类路径。

    不需要修改主控文件，也不需要额外传入model_name。
    """
    names = []

    names.append(net.__class__.__name__.lower())
    names.append(
        f'{net.__class__.__module__}.{net.__class__.__name__}'.lower()
    )

    inner_model = getattr(net, 'model', None)
    if inner_model is not None:
        names.append(inner_model.__class__.__name__.lower())
        names.append(
            f'{inner_model.__class__.__module__}.'
            f'{inner_model.__class__.__name__}'.lower()
        )

    for module in net.modules():
        names.append(module.__class__.__name__.lower())
        names.append(
            f'{module.__class__.__module__}.'
            f'{module.__class__.__name__}'.lower()
        )

    return any(
        model_name in name
        for name in names
        for model_name in ADVANCED_MODEL_NAMES
    )


def _dataset_elements(dataset):
    """
    统一tuple数据集和tuple列表的遍历方式。
    """
    if type(dataset) is tuple:
        return [dataset]
    return dataset


def _segmented_forward_backward(
    device,
    dataset,
    net,
    scaler,
    segment_length
):
    """
    Transformer/Informer/Autoformer的分段前向和反向。

    每段执行forward和backward，全部分段完成后由net_train统一执行一次
    optimizer.step()，因此不是每段更新一次参数。

    返回：
        outputs
        output_dataset
    """
    outputs_all = []
    output_dataset_all = []

    dataset_elements = _dataset_elements(dataset)

    total_output_elements = 0
    for _, output_dataset in dataset_elements:
        total_output_elements += int(output_dataset.numel())

    if total_output_elements <= 0:
        raise ValueError('output_dataset为空，无法训练。')

    for input_dataset, output_dataset in dataset_elements:
        input_dataset = input_dataset.to(device)
        output_dataset = output_dataset.to(device)

        sequence_length, input_dim = input_dataset.shape

        for start in range(0, sequence_length, segment_length):
            end = min(start + segment_length, sequence_length)

            # 避免最后仅剩1个时间步时，部分时序模型内部长度计算失效。
            if end - start < 2:
                continue

            input_segment = input_dataset[start:end]
            output_segment = output_dataset[start:end]

            segment_steps = input_segment.shape[0]
            input_segment = input_segment.reshape(
                segment_steps,
                1,
                input_dim
            )

            # 各段MSE按输出元素数量加权。
            # 所有分段梯度累计后，与全序列mean MSE的梯度尺度保持一致。
            loss_weight = (
                output_segment.numel()
                / total_output_elements
            )

            if scaler is not None:
                with amp.autocast():
                    outputs_segment = net.forward(
                        input_segment
                    ).reshape(output_segment.shape)

                    loss_segment = F.mse_loss(
                        outputs_segment,
                        output_segment
                    ) * loss_weight

                scaler.scale(loss_segment).backward()
            else:
                outputs_segment = net.forward(
                    input_segment
                ).reshape(output_segment.shape)

                loss_segment = F.mse_loss(
                    outputs_segment,
                    output_segment
                ) * loss_weight

                loss_segment.backward()

            # detach防止保存全部分段计算图，降低显存占用。
            outputs_all.append(outputs_segment.detach())
            output_dataset_all.append(output_segment.detach())

    if len(outputs_all) == 0:
        raise ValueError(
            '没有生成有效分段，请检查数据长度和'
            'ADVANCED_MODEL_SEGMENT_LENGTH。'
        )

    outputs = torch.cat(outputs_all, dim=0)
    output_dataset = torch.cat(output_dataset_all, dim=0)

    return outputs, output_dataset


def _segmented_forward(
    device,
    dataset,
    net,
    segment_length
):
    """
    Transformer/Informer/Autoformer的分段测试前向。
    """
    outputs_all = []
    output_dataset_all = []

    for input_dataset, output_dataset in _dataset_elements(dataset):
        input_dataset = input_dataset.to(device)
        output_dataset = output_dataset.to(device)

        sequence_length, input_dim = input_dataset.shape

        for start in range(0, sequence_length, segment_length):
            end = min(start + segment_length, sequence_length)

            if end - start < 2:
                continue

            input_segment = input_dataset[start:end]
            output_segment = output_dataset[start:end]

            segment_steps = input_segment.shape[0]
            input_segment = input_segment.reshape(
                segment_steps,
                1,
                input_dim
            )

            outputs_segment = net.forward(
                input_segment
            ).reshape(output_segment.shape)

            outputs_all.append(outputs_segment.detach())
            output_dataset_all.append(output_segment.detach())

    if len(outputs_all) == 0:
        raise ValueError(
            '没有生成有效分段，请检查数据长度和'
            'ADVANCED_MODEL_SEGMENT_LENGTH。'
        )

    outputs = torch.cat(outputs_all, dim=0)
    output_dataset = torch.cat(output_dataset_all, dim=0)

    return outputs, output_dataset


# 训练函数
# 函数名称和参数接口保持不变
def net_train(
    device,
    dataset,
    optimizer,
    net,
    scaler,
    out_dir=None,
    iteration=None,
    draw_pic=False,
    debug_grad=False,
    debug_net=None
):
    optimizer.zero_grad()

    do_debug = False  # 检测梯度
    do_print_grad = True  # 查看梯度
    do_grad_clip = True  # 梯度裁剪
    detect_grad_nan = True  # 检测梯度是否为NAN,并剔除

    advanced_attention_model = _is_advanced_attention_model(net)
    gradient_net = debug_net if debug_net is not None else net

    if advanced_attention_model:
        # Transformer、Informer、Autoformer：
        # 在原net_train内部自动执行分段计算，主控接口无需修改。
        outputs, output_dataset = _segmented_forward_backward(
            device=device,
            dataset=dataset,
            net=net,
            scaler=scaler,
            segment_length=ADVANCED_MODEL_SEGMENT_LENGTH
        )

        if scaler is not None:
            # 将缩放后的梯度恢复为真实梯度，再计算L2范数。
            scaler.unscale_(optimizer)
            gradient_l2_norm = global_gradient_l2_norm(gradient_net)
            scaler.step(optimizer)
            scaler.update()
        else:
            gradient_l2_norm = global_gradient_l2_norm(gradient_net)
            optimizer.step()

    elif scaler is not None:
        # 其他模型保持原来的整段AMP训练逻辑。
        with amp.autocast():
            loss, outputs, output_dataset = get_loss(
                dataset=dataset,
                device=device,
                net=net
            )

        if debug_grad:
            lebw.test_grad_error(
                net=gradient_net,
                loss=loss,
                do_debug=do_debug,
                do_print_grad=do_print_grad,
                do_grad_clip=do_grad_clip,
                detect_grad_nan=detect_grad_nan,
                out_dir=out_dir,
                optimizer=optimizer,
                iteration=iteration,
                scaler=scaler
            )
        else:
            scaler.scale(loss).backward()
            scaler.unscale_(optimizer)

        gradient_l2_norm = global_gradient_l2_norm(gradient_net)

        scaler.step(optimizer)
        scaler.update()

    else:
        # 其他模型保持原来的整段非AMP训练逻辑。
        loss, outputs, output_dataset = get_loss(
            dataset=dataset,
            device=device,
            net=net
        )

        if debug_grad:
            lebw.test_grad_error(
                net=gradient_net,
                loss=loss,
                do_debug=do_debug,
                do_print_grad=do_print_grad,
                do_grad_clip=do_grad_clip,
                detect_grad_nan=detect_grad_nan,
                out_dir=out_dir,
                optimizer=optimizer,
                iteration=iteration
            )
        else:
            loss.backward()

        gradient_l2_norm = global_gradient_l2_norm(gradient_net)
        optimizer.step()

    nrmse, rmse, nmse, mse, mae = narma_estimate.nrmse(
        outputs,
        output_dataset
    )

    if draw_pic:
        tp.draw_lorenz(
            predict=outputs.cpu().detach(),
            true=output_dataset.cpu().detach()
        )

    return nrmse, rmse, nmse, mse, mae, gradient_l2_norm


# 测试函数
# 函数名称和参数接口保持不变
def net_test(device, dataset, net, draw_pic=False):
    with torch.no_grad():
        if _is_advanced_attention_model(net):
            # Transformer、Informer、Autoformer自动分段测试。
            outputs, output_dataset = _segmented_forward(
                device=device,
                dataset=dataset,
                net=net,
                segment_length=ADVANCED_MODEL_SEGMENT_LENGTH
            )
        else:
            # 其他模型保持原来的整段测试逻辑。
            loss, outputs, output_dataset = get_loss(
                dataset=dataset,
                device=device,
                net=net
            )

    nrmse, rmse, nmse, mse, mae = narma_estimate.nrmse(
        outputs,
        output_dataset
    )

    if draw_pic:
        tp.draw_lorenz(
            predict=outputs.cpu().detach(),
            true=output_dataset.cpu().detach()
        )

    return nrmse, rmse, nmse, mse, mae


def main():
    # 加载中断文件
    # resume = f'./linear_amp/t(200)_sgd_lr0.01/checkpoint_max.pth'
    resume = ''

    node_num = 90  # 老传统90个隐藏层节点
    device = torch.device(
        'cuda:0' if torch.cuda.is_available() else 'cpu'
    )

    random_bptt = False

    dataset_names = ['ETTh1', 'Lorenz', 'NARMA']
    name = dataset_names[2]

    input_dim, output_dim, train_set, test_set, predict_step = dataset_from(
        name=name,
        device=device,
        random_bptt=random_bptt
    )

    # 定义模型
    device = torch.device(
        'cuda:0' if torch.cuda.is_available() else 'cpu'
    )
    net = three_linear_net(
        in_dim=input_dim,
        n_hidden=node_num,
        out_dim=output_dim
    ).to(device)

    # 定义损失函数和优化器
    opt = 'sgd'
    lr = 1e-2
    momentum = 0.9

    if opt == 'sgd':
        optimizer = torch.optim.SGD(
            net.parameters(),
            lr=lr,
            momentum=momentum
        )
    elif opt == 'adam':
        optimizer = torch.optim.Adam(
            net.parameters(),
            lr=lr
        )
    else:
        raise NotImplementedError(opt)

    # 定义超参数
    max_iteration = 200
    use_amp = True

    lr_scheduler = torch.optim.lr_scheduler.CosineAnnealingLR(
        optimizer,
        max_iteration
    )

    # 设置文件输出位置
    out_dir = './linear'
    scaler = None

    if use_amp:
        out_dir += '_amp'
        scaler = amp.GradScaler()

    out_dir = os.path.join(
        out_dir,
        f't({max_iteration})_{name}_lr{lr}'
    )

    if not os.path.exists(out_dir):
        os.makedirs(out_dir)
        print(f'Mkdir {out_dir}.')
    else:
        print(f'EXdir {out_dir}.')

    start_epoch = 0
    min_test_nrmse = 100

    if resume:
        checkpoint = torch.load(
            resume,
            map_location='cpu'
        )
        net.load_state_dict(checkpoint['net'])
        optimizer.load_state_dict(checkpoint['optimizer'])
        lr_scheduler.load_state_dict(
            checkpoint['lr_scheduler']
        )
        start_epoch = checkpoint['iteration'] + 1
        min_test_nrmse = checkpoint['min_test_nrmse']

    result_len = 5
    train_result_list = np.zeros(
        shape=[max_iteration, result_len]
    )
    test_result_list = np.zeros(
        shape=[max_iteration, result_len]
    )

    for iteration in range(start_epoch, max_iteration):
        start_time = time.time()

        if iteration in [50, 100, 150, 199]:
            draw_pic = True
        else:
            draw_pic = False

        net.train()
        train_result = net_train(
            device=device,
            dataset=train_set,
            optimizer=optimizer,
            net=net,
            scaler=scaler,
            out_dir=out_dir,
            iteration=iteration,
            draw_pic=draw_pic,
            debug_grad=False
        )

        train_nrmse = train_result[0]

        # net_train返回5个性能指标和1个梯度L2；
        # 原main中的性能结果表继续保存前5项。
        train_result_list[iteration] = np.array(
            train_result[:5]
        )

        lr_scheduler.step()

        net.eval()
        test_result = net_test(
            device=device,
            dataset=test_set,
            net=net,
            draw_pic=draw_pic
        )

        test_nrmse = test_result[0]
        test_result_list[iteration] = np.array(test_result)

        save_max = False

        if test_nrmse < min_test_nrmse:
            min_test_nrmse = test_nrmse
            save_max = True

        checkpoint = {
            'net': net.state_dict(),
            'optimizer': optimizer.state_dict(),
            'lr_scheduler': lr_scheduler.state_dict(),
            'iteration': iteration,
            'min_test_nrmse': min_test_nrmse
        }

        if save_max:
            torch.save(
                checkpoint,
                os.path.join(out_dir, 'checkpoint_max.pth')
            )

        torch.save(
            checkpoint,
            os.path.join(out_dir, 'checkpoint_latest.pth')
        )

        wb = openpyxl.Workbook()
        train_excel = wb.create_sheet('train')
        test_excel = wb.create_sheet('test')

        for evalu_para in range(train_result_list.shape[1]):
            train_excel.append(
                train_result_list[:, evalu_para].tolist()
            )
            test_excel.append(
                test_result_list[:, evalu_para].tolist()
            )

        file_name = (
            f'./ETT_{predict_step}_'
            f'({datetime.date.today()}).xlsx'
        )

        wb.save(os.path.join(out_dir, file_name))

        print(out_dir)
        print(
            f'iteration = {iteration}, '
            f'train_nrmse ={train_nrmse: .4f}, '
            f'test_nrmse ={test_nrmse: .4f}, '
            f'min_test_nrmse ={min_test_nrmse: .4f}'
        )
        print(
            f'escape time = '
            f'{(datetime.datetime.now() + datetime.timedelta(seconds=(time.time() - start_time) * (max_iteration - iteration))).strftime("%Y-%m-%d %H:%M:%S")}\n'
        )


if __name__ == '__main__':
    main()
